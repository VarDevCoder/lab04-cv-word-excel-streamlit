"""
Laboratorio 4: Interfaz Gráfica de Escritorio (Tkinter) para Ingestión de CVs
Word (.docx) y Homologación de Excel (.xlsx)

Cátedra: Aplicaciones de la Inteligencia Artificial
Licenciatura en Análisis de Sistemas - UNIBE

Versión de escritorio equivalente a app.py (Streamlit), construida solo con
la librería estándar de Python (tkinter) más python-docx / pandas / openpyxl.
"""

import io
import re
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import pandas as pd
from docx import Document

ESQUEMA_ESTANDAR = {
    "nombre": ["nombre", "nombre completo", "apellido y nombre", "candidato", "name"],
    "email": ["email", "correo", "correo electronico", "e-mail", "mail"],
    "telefono": ["telefono", "teléfono", "celular", "phone", "contacto"],
    "cargo": ["cargo", "puesto", "posicion", "posición", "role", "position"],
    "experiencia_anios": [
        "experiencia",
        "años experiencia",
        "anios experiencia",
        "experience",
        "años",
    ],
}


def normalizar(texto: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(texto).strip().lower())


def sugerir_columna(col_original: str) -> str | None:
    col_norm = normalizar(col_original)
    for campo_estandar, alias in ESQUEMA_ESTANDAR.items():
        for a in alias:
            if normalizar(a) == col_norm:
                return campo_estandar
    return None


class CVTab(ttk.Frame):
    """Pestaña de ingestión y procesamiento de CVs en Word."""

    def __init__(self, parent: ttk.Notebook):
        super().__init__(parent, padding=12)
        self.parrafos: list[str] = []
        self.doc = None
        self._build_ui()

    def _build_ui(self):
        top = ttk.Frame(self)
        top.pack(fill="x")
        ttk.Button(top, text="📄 Seleccionar CV (.docx)", command=self.cargar_docx).pack(
            side="left"
        )
        self.lbl_resumen = ttk.Label(top, text="Ningún archivo cargado.")
        self.lbl_resumen.pack(side="left", padx=12)

        info = ttk.Frame(self)
        info.pack(fill="x", pady=8)
        self.lbl_email = ttk.Label(info, text="📧 Correo: —")
        self.lbl_email.pack(side="left", padx=(0, 20))
        self.lbl_telefono = ttk.Label(info, text="📱 Teléfono: —")
        self.lbl_telefono.pack(side="left")

        notebook_interno = ttk.Notebook(self)
        notebook_interno.pack(fill="both", expand=True, pady=(8, 0))

        # Sub-tab: Ficha formateada
        self.tab_ficha = ttk.Frame(notebook_interno, padding=8)
        notebook_interno.add(self.tab_ficha, text="📋 Ficha Formateada")
        self.txt_ficha = tk.Text(self.tab_ficha, wrap="word")
        self.txt_ficha.pack(fill="both", expand=True)
        self.txt_ficha.configure(state="disabled")

        # Sub-tab: Texto editable y descarga
        self.tab_editable = ttk.Frame(notebook_interno, padding=8)
        notebook_interno.add(self.tab_editable, text="✏️ Texto Editable & Descarga")
        self.txt_editable = tk.Text(self.tab_editable, wrap="word")
        self.txt_editable.pack(fill="both", expand=True)
        ttk.Button(
            self.tab_editable,
            text="💾 Descargar CV Procesado (.txt)",
            command=self.descargar_txt,
        ).pack(pady=8)

        # Sub-tab: Tablas
        self.tab_tablas = ttk.Frame(notebook_interno, padding=8)
        notebook_interno.add(self.tab_tablas, text="📊 Tablas del Documento")
        self.tree_tablas = ttk.Treeview(self.tab_tablas, show="headings")
        self.tree_tablas.pack(fill="both", expand=True)

    def cargar_docx(self):
        ruta = filedialog.askopenfilename(
            title="Selecciona un CV en formato Word",
            filetypes=[("Documentos Word", "*.docx")],
        )
        if not ruta:
            return

        with open(ruta, "rb") as f:
            doc_bytes = io.BytesIO(f.read())
        self.doc = Document(doc_bytes)
        self.parrafos = [p.text.strip() for p in self.doc.paragraphs if p.text.strip()]

        self.lbl_resumen.config(
            text=f"✅ {len(self.parrafos)} párrafos y {len(self.doc.tables)} tablas encontradas."
        )

        texto_completo = "\n".join(self.parrafos)
        email_match = re.search(
            r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", texto_completo
        )
        telefono_match = re.search(
            r"(\+?\d{1,4}[\s-]?)?(\(?\d{2,4}\)?[\s-]?)?\d{3,4}[\s-]?\d{3,4}",
            texto_completo,
        )
        self.lbl_email.config(
            text=f"📧 Correo: {email_match.group(0) if email_match else 'No detectado'}"
        )
        self.lbl_telefono.config(
            text=f"📱 Teléfono: {telefono_match.group(0) if telefono_match else 'No detectado'}"
        )

        self._poblar_ficha()
        self._poblar_editable(texto_completo)
        self._poblar_tablas()

    def _poblar_ficha(self):
        self.txt_ficha.configure(state="normal")
        self.txt_ficha.delete("1.0", tk.END)
        for p in self.parrafos[:6]:
            if p.isupper():
                self.txt_ficha.insert(tk.END, f"📌 {p}\n", "titulo")
            else:
                self.txt_ficha.insert(tk.END, f"• {p}\n")
        self.txt_ficha.tag_configure("titulo", font=("TkDefaultFont", 11, "bold"))
        self.txt_ficha.configure(state="disabled")

    def _poblar_editable(self, texto_completo: str):
        self.txt_editable.delete("1.0", tk.END)
        self.txt_editable.insert(tk.END, texto_completo)

    def _poblar_tablas(self):
        self.tree_tablas.delete(*self.tree_tablas.get_children())
        self.tree_tablas["columns"] = ()

        if not self.doc.tables:
            return

        tabla = self.doc.tables[0]
        filas = [[celda.text.strip() for celda in fila.cells] for fila in tabla.rows]
        if len(filas) <= 1:
            return

        columnas = filas[0]
        self.tree_tablas["columns"] = columnas
        for col in columnas:
            self.tree_tablas.heading(col, text=col)
            self.tree_tablas.column(col, width=120)
        for fila in filas[1:]:
            self.tree_tablas.insert("", tk.END, values=fila)

    def descargar_txt(self):
        contenido = self.txt_editable.get("1.0", tk.END).strip()
        if not contenido:
            messagebox.showwarning("Aviso", "No hay texto para descargar.")
            return
        ruta = filedialog.asksaveasfilename(
            title="Guardar CV procesado",
            defaultextension=".txt",
            initialfile="cv_procesado.txt",
            filetypes=[("Archivo de texto", "*.txt")],
        )
        if not ruta:
            return
        with open(ruta, "w", encoding="utf-8") as f:
            f.write(contenido)
        messagebox.showinfo("Éxito", f"Archivo guardado en:\n{ruta}")


class ExcelTab(ttk.Frame):
    """Pestaña de mapeo y homologación de columnas Excel."""

    def __init__(self, parent: ttk.Notebook):
        super().__init__(parent, padding=12)
        self.df_original: pd.DataFrame | None = None
        self.combo_vars: dict[str, tk.StringVar] = {}
        self._build_ui()

    def _build_ui(self):
        top = ttk.Frame(self)
        top.pack(fill="x")
        ttk.Button(
            top, text="📊 Seleccionar Excel (.xlsx)", command=self.cargar_excel
        ).pack(side="left")
        self.lbl_resumen = ttk.Label(top, text="Ningún archivo cargado.")
        self.lbl_resumen.pack(side="left", padx=12)

        ttk.Label(
            self, text="Vista previa de los datos originales:", font=("TkDefaultFont", 10, "bold")
        ).pack(anchor="w", pady=(12, 4))
        self.tree_preview = ttk.Treeview(self, show="headings", height=8)
        self.tree_preview.pack(fill="both", expand=True)

        ttk.Label(
            self, text="Mapeo de columnas hacia el esquema estándar:",
            font=("TkDefaultFont", 10, "bold"),
        ).pack(anchor="w", pady=(12, 4))

        self.frame_mapeo = ttk.Frame(self)
        self.frame_mapeo.pack(fill="x")

        ttk.Button(
            self,
            text="🔄 Homologar y Guardar Planilla Estándar (.xlsx)",
            command=self.homologar_y_guardar,
        ).pack(pady=12)

    def cargar_excel(self):
        ruta = filedialog.askopenfilename(
            title="Selecciona una planilla Excel",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not ruta:
            return

        self.df_original = pd.read_excel(ruta, engine="openpyxl")
        self.lbl_resumen.config(
            text=f"✅ {len(self.df_original)} filas, {len(self.df_original.columns)} columnas."
        )

        self._poblar_preview()
        self._poblar_mapeo()

    def _poblar_preview(self):
        self.tree_preview.delete(*self.tree_preview.get_children())
        columnas = list(self.df_original.columns)
        self.tree_preview["columns"] = columnas
        for col in columnas:
            self.tree_preview.heading(col, text=col)
            self.tree_preview.column(col, width=120)
        for _, fila in self.df_original.head(50).iterrows():
            self.tree_preview.insert("", tk.END, values=list(fila))

    def _poblar_mapeo(self):
        for widget in self.frame_mapeo.winfo_children():
            widget.destroy()
        self.combo_vars.clear()

        opciones = ["(Ignorar columna)"] + list(ESQUEMA_ESTANDAR.keys())

        for idx, columna in enumerate(self.df_original.columns):
            ttk.Label(self.frame_mapeo, text=f"«{columna}»  →").grid(
                row=idx, column=0, sticky="w", padx=(0, 8), pady=2
            )
            sugerencia = sugerir_columna(columna) or "(Ignorar columna)"
            var = tk.StringVar(value=sugerencia)
            combo = ttk.Combobox(
                self.frame_mapeo, textvariable=var, values=opciones, state="readonly", width=22
            )
            combo.grid(row=idx, column=1, sticky="w", pady=2)
            self.combo_vars[columna] = var

    def homologar_y_guardar(self):
        if self.df_original is None:
            messagebox.showwarning("Aviso", "Primero carga una planilla Excel.")
            return

        mapeo = {
            col: var.get()
            for col, var in self.combo_vars.items()
            if var.get() != "(Ignorar columna)"
        }
        if not mapeo:
            messagebox.showwarning("Aviso", "Selecciona al menos una columna para homologar.")
            return

        df_homologado = self.df_original[list(mapeo.keys())].rename(columns=mapeo)

        ruta = filedialog.asksaveasfilename(
            title="Guardar planilla homologada",
            defaultextension=".xlsx",
            initialfile="planilla_homologada.xlsx",
            filetypes=[("Excel", "*.xlsx")],
        )
        if not ruta:
            return

        df_homologado.to_excel(ruta, index=False, engine="openpyxl")
        messagebox.showinfo("Éxito", f"Planilla homologada guardada en:\n{ruta}")


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Lab 04 — Procesador de CVs (Word) & Homologador de Excel")
        self.geometry("900x650")

        notebook = ttk.Notebook(self)
        notebook.pack(fill="both", expand=True)

        notebook.add(CVTab(notebook), text="📄 Procesador de CVs (Word)")
        notebook.add(ExcelTab(notebook), text="📊 Homologador de Excel")


if __name__ == "__main__":
    App().mainloop()
