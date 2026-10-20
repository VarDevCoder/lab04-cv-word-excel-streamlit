"""
Laboratorio 4: App Web en Python (Streamlit) para Ingestión de CVs Word (.docx)
y Homologación de Excel (.xlsx)

Cátedra: Aplicaciones de la Inteligencia Artificial
Licenciatura en Análisis de Sistemas - UNIBE
"""

import io
import re

import pandas as pd
import streamlit as st
from docx import Document

st.set_page_config(
    page_title="Lab 04 · CVs & Excel",
    page_icon="🧠",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ---------------------------------------------------------------------------
# Estilos (tema custom)
# ---------------------------------------------------------------------------
st.markdown(
    """
    <style>
        :root {
            --accent: #7C5CFC;
            --accent-dark: #5B3FD6;
            --bg-card: #1A1C29;
        }

        .stApp {
            background: radial-gradient(circle at 15% 0%, #241f3d 0%, #0E0F1A 45%);
        }

        #MainMenu, footer, header[data-testid="stHeader"] { visibility: hidden; }

        .block-container { padding-top: 2rem; max-width: 1150px; }

        .hero {
            background: linear-gradient(120deg, #7C5CFC 0%, #4E3DB8 100%);
            border-radius: 20px;
            padding: 2rem 2.4rem;
            margin-bottom: 1.6rem;
            box-shadow: 0 18px 40px -16px rgba(124, 92, 252, 0.55);
        }
        .hero h1 {
            color: white;
            font-size: 2.1rem;
            margin: 0 0 .35rem 0;
            font-weight: 800;
        }
        .hero p {
            color: rgba(255,255,255,0.85);
            margin: 0;
            font-size: 1.02rem;
        }
        .hero .badge {
            display: inline-block;
            background: rgba(255,255,255,0.18);
            color: white;
            padding: .25rem .7rem;
            border-radius: 999px;
            font-size: .78rem;
            font-weight: 600;
            letter-spacing: .03em;
            margin-bottom: .6rem;
        }

        div[data-testid="stMetric"] {
            background: var(--bg-card);
            border: 1px solid rgba(124, 92, 252, 0.25);
            border-radius: 14px;
            padding: 1rem 1.2rem;
        }
        div[data-testid="stMetric"] label { color: #B7B4CF !important; }

        .card {
            background: var(--bg-card);
            border: 1px solid rgba(255,255,255,0.06);
            border-radius: 16px;
            padding: 1.4rem 1.6rem;
            margin-bottom: 1.2rem;
        }
        .card h4 { margin-top: 0; color: #E8E6F7; }

        .pill {
            display: inline-block;
            background: rgba(124, 92, 252, 0.15);
            color: #B9A9FF;
            border: 1px solid rgba(124, 92, 252, 0.35);
            padding: .15rem .6rem;
            border-radius: 999px;
            font-size: .75rem;
            font-weight: 600;
            margin-right: .4rem;
        }

        .stTabs [data-baseweb="tab-list"] { gap: 6px; }
        .stTabs [data-baseweb="tab"] {
            background: var(--bg-card);
            border-radius: 10px 10px 0 0;
            padding: 10px 18px;
            color: #B7B4CF;
        }
        .stTabs [aria-selected="true"] {
            background: var(--accent) !important;
            color: white !important;
        }
        .stTabs [data-baseweb="tab-highlight"] { background-color: var(--accent) !important; }
        .stTabs [data-baseweb="tab-border"] { background-color: transparent !important; }

        .stButton > button, .stDownloadButton > button {
            background: linear-gradient(120deg, var(--accent), var(--accent-dark));
            color: white;
            border: none;
            border-radius: 10px;
            font-weight: 600;
            padding: .55rem 1.2rem;
        }
        .stButton > button:hover, .stDownloadButton > button:hover {
            filter: brightness(1.1);
            color: white;
        }

        section[data-testid="stFileUploaderDropzone"] {
            background: var(--bg-card);
            border: 1.5px dashed rgba(124, 92, 252, 0.45);
            border-radius: 14px;
        }
    </style>
    """,
    unsafe_allow_html=True,
)

# ---------------------------------------------------------------------------
# Hero / Header
# ---------------------------------------------------------------------------
st.markdown(
    """
    <div class="hero">
        <span class="badge">LAB 04 · UNIDAD II · CÁTEDRA IA</span>
        <h1>🧠 Procesador de CVs &amp; Homologador de Excel</h1>
        <p>Ingestión de currículums en Word (.docx) en memoria y homologación
        inteligente de planillas Excel (.xlsx) con Streamlit + python-docx + pandas.</p>
    </div>
    """,
    unsafe_allow_html=True,
)

tab_cv, tab_excel = st.tabs(["📄  Procesador de CVs (Word)", "📊  Homologador de Excel"])

# ---------------------------------------------------------------------------
# TAB 1 — Ingestión y procesamiento de CVs en Word (.docx)
# ---------------------------------------------------------------------------
with tab_cv:
    st.markdown(
        """
        <div class="card">
            <h4>📄 Sube un currículum en Word</h4>
            <span class="pill">io.BytesIO</span>
            <span class="pill">python-docx</span>
            <span class="pill">regex</span>
        </div>
        """,
        unsafe_allow_html=True,
    )

    archivo_subido = st.file_uploader(
        "Selecciona un CV en formato Word (.docx)", type=["docx"]
    )

    if archivo_subido:
        # Paso 1: Ingestión directa en memoria (segura, sin archivos temporales en disco)
        doc_bytes = io.BytesIO(archivo_subido.read())
        doc = Document(doc_bytes)
        parrafos = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

        st.success(
            f"✅ Documento procesado: **{len(parrafos)}** párrafos y "
            f"**{len(doc.tables)}** tablas encontradas."
        )

        # Paso 2: Extracción heurística de entidades (regex) y tablas estructuradas
        texto_completo = "\n".join(parrafos)

        email_match = re.search(
            r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", texto_completo
        )
        telefono_match = re.search(
            r"(\+?\d{1,4}[\s-]?)?(\(?\d{2,4}\)?[\s-]?)?\d{3,4}[\s-]?\d{3,4}",
            texto_completo,
        )

        col1, col2, col3 = st.columns(3)
        with col1:
            st.metric(
                "📧 Correo Detectado",
                email_match.group(0) if email_match else "No detectado",
            )
        with col2:
            st.metric(
                "📱 Teléfono Detectado",
                telefono_match.group(0) if telefono_match else "No detectado",
            )
        with col3:
            st.metric("🧾 Párrafos extraídos", len(parrafos))

        if doc.tables:
            st.markdown("#### 📊 Tablas Encontradas en el Documento")
            for idx, tabla in enumerate(doc.tables):
                filas = [
                    [celda.text.strip() for celda in fila.cells] for fila in tabla.rows
                ]
                if len(filas) > 1:
                    df_tabla = pd.DataFrame(filas[1:], columns=filas[0])
                    st.caption(f"Tabla {idx + 1}")
                    st.dataframe(df_tabla, use_container_width=True)

        # Paso 3: Ficha de candidato, texto editable y exportación
        st.divider()
        tab1, tab2 = st.tabs(["📋 Ficha Formateada", "✏️ Texto Editable & Descarga"])

        with tab1:
            st.markdown("##### Perfil del Postulante")
            for p in parrafos[:6]:
                if p.isupper():
                    st.markdown(f"**📌 {p}**")
                else:
                    st.write(f"• {p}")

        with tab2:
            texto_editado = st.text_area(
                "Texto extraído (editable para notas de entrevista):",
                value=texto_completo,
                height=250,
            )
            st.download_button(
                label="💾 Descargar CV Procesado (.txt)",
                data=texto_editado,
                file_name="cv_procesado.txt",
                mime="text/plain",
            )
    else:
        st.info("⬆️ Sube un archivo .docx para comenzar el procesamiento.")

# ---------------------------------------------------------------------------
# TAB 2 — Desafío de extensión: Homologador de columnas Excel
# ---------------------------------------------------------------------------
with tab_excel:
    st.markdown(
        """
        <div class="card">
            <h4>📊 Mapeador y Homologador de Columnas Excel</h4>
            <span class="pill">pandas</span>
            <span class="pill">openpyxl</span>
            <span class="pill">auto-mapeo</span>
            <p style="margin:.6rem 0 0 0; color:#B7B4CF;">
                Sube una planilla con nombres de columna desordenados o inconsistentes
                y homológala a un esquema estándar antes de descargarla.
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    ESQUEMA_ESTANDAR = {
        "nombre": ["nombre", "nombre completo", "apellido y nombre", "candidato", "name"],
        "email": ["email", "correo", "correo electronico", "e-mail", "mail"],
        "telefono": ["telefono", "teléfono", "celular", "phone", "contacto"],
        "cargo": ["cargo", "puesto", "posicion", "posición", "role", "position"],
        "experiencia_anios": [
            "experiencia",
            "años experiencia",
            "anios experiencia",
            "experience",
            "años",
        ],
    }

    def normalizar(texto: str) -> str:
        return re.sub(r"[^a-z0-9]", "", str(texto).strip().lower())

    def sugerir_columna(col_original: str) -> str | None:
        col_norm = normalizar(col_original)
        for campo_estandar, alias in ESQUEMA_ESTANDAR.items():
            for a in alias:
                if normalizar(a) == col_norm:
                    return campo_estandar
        return None

    archivo_excel = st.file_uploader(
        "Selecciona una planilla Excel (.xlsx)", type=["xlsx"], key="excel_uploader"
    )

    if archivo_excel:
        excel_bytes = io.BytesIO(archivo_excel.read())
        df_original = pd.read_excel(excel_bytes, engine="openpyxl")

        m1, m2 = st.columns(2)
        m1.metric("📑 Filas", len(df_original))
        m2.metric("🧮 Columnas", len(df_original.columns))

        st.markdown("##### Vista previa de los datos originales")
        st.dataframe(df_original, use_container_width=True)

        st.divider()
        st.markdown("##### 🔗 Mapeo de columnas")
        st.caption(
            "Revisa las sugerencias automáticas y ajusta manualmente las columnas "
            "que no coincidan con el esquema estándar."
        )

        opciones_destino = ["(Ignorar columna)"] + list(ESQUEMA_ESTANDAR.keys())
        mapeo = {}

        cols_grid = st.columns(2)
        for i, columna in enumerate(df_original.columns):
            sugerencia = sugerir_columna(columna)
            indice_default = (
                opciones_destino.index(sugerencia) if sugerencia in opciones_destino else 0
            )
            with cols_grid[i % 2]:
                seleccion = st.selectbox(
                    f"Columna original: «{columna}»",
                    options=opciones_destino,
                    index=indice_default,
                    key=f"map_{columna}",
                )
            if seleccion != "(Ignorar columna)":
                mapeo[columna] = seleccion

        st.write("")
        if st.button("🔄 Homologar y Generar Planilla Estándar"):
            if mapeo:
                df_homologado = df_original[list(mapeo.keys())].rename(columns=mapeo)
                st.success("✅ Planilla homologada correctamente.")
                st.dataframe(df_homologado, use_container_width=True)

                buffer_salida = io.BytesIO()
                df_homologado.to_excel(buffer_salida, index=False, engine="openpyxl")
                buffer_salida.seek(0)

                st.download_button(
                    label="💾 Descargar Excel Homologado (.xlsx)",
                    data=buffer_salida,
                    file_name="planilla_homologada.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            else:
                st.warning("⚠️ Selecciona al menos una columna para homologar.")
    else:
        st.info("⬆️ Sube un archivo .xlsx para comenzar la homologación.")

st.markdown(
    """
    <div style="text-align:center; color:#5C5A78; font-size:.8rem; margin-top:2.5rem;">
        Cátedra Aplicaciones de la Inteligencia Artificial · Licenciatura en Análisis de Sistemas · UNIBE 2026
    </div>
    """,
    unsafe_allow_html=True,
)
